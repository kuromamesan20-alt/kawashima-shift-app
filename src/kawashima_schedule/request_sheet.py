"""希望休・希望出勤の入力シート(月ごと)の書き出しと読み戻し。

施設の担当者が毎月これに記入する。31名×日付のマス目で、セルはプルダウン。
「公」を選べば希望休、シフトの記号を選べばその勤務での希望出勤になる。

シート自体は毎月こちらで自動生成するので、施設の担当者は同じ場所を開くだけでよい。
"""

from __future__ import annotations

import calendar
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Sequence, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .calendar_utils import Day, month_days, month_label
from .models import StaffProfile
from .shifts import (
    ABSENCE_MARKS,
    DAY_SHIFTS,
    HOUR_CHOICES,
    LATE_NIGHT_IN,
    NIGHT_AFTER,
    NIGHT_IN,
    OFF,
)

SHEET_NAME = "希望"
LEGEND_SHEET = "書き方"

# セルのプルダウンに出す選択肢。空欄=希望なし。
WISH_OFF = OFF  # 公 = この日は休みたい
WISH_CHOICES: Tuple[str, ...] = (
    (WISH_OFF,)
    + tuple(DAY_SHIFTS)
    + HOUR_CHOICES  # 時短。番号のシフトに当てはまらない時間帯で働く人用
    + (NIGHT_IN, LATE_NIGHT_IN)
    + ABSENCE_MARKS
)

# 前月の最終日の記号を「0日」として持つ。
# 夜勤は ○→△→公、深夜は ◉→公 と複数日にまたがるので、
# 前月末が分からないと今月1日・2日が前月と食い違う。
# 実物の勤務表でも、毎月1日には前月から続く △ が2人いる。
CARRY_OVER_DAY = 0
CARRY_OVER_CHOICES: Tuple[str, ...] = (NIGHT_IN, NIGHT_AFTER, LATE_NIGHT_IN)

_ID_COLUMN = 1
_NAME_COLUMN = 2
_CARRY_COLUMN = 3  # 前月の最終日
_FIRST_DAY_COLUMN = 4
_HEADER_ROW = 3

_HEAD_FILL = PatternFill("solid", fgColor="2F5D8C")
_WEEKEND_FILL = PatternFill("solid", fgColor="FDECEC")
_NAME_FILL = PatternFill("solid", fgColor="EDF0F4")
_OFF_FILL = PatternFill("solid", fgColor="EFEFEF")
_THIN = Side(style="thin", color="D9DEE5")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


class RequestSheetError(ValueError):
    """入力シートの形式が想定と違う場合に送出する。"""


@dataclass
class StaffRequests:
    """スタッフ1人分の、その月の希望。"""

    staff_id: str
    name: str
    # 日にち(1〜31) -> 記号。「公」なら希望休、シフト記号なら希望出勤。
    entries: Dict[int, str] = field(default_factory=dict)

    @property
    def carry_over(self) -> str:
        """前月の最終日の記号(○ / △ / ◉)。無ければ空。"""
        return self.entries.get(CARRY_OVER_DAY, "")

    def wish_off_days(self) -> List[int]:
        return sorted(
            day
            for day, mark in self.entries.items()
            if mark == WISH_OFF and day >= 1
        )

    def wish_work_days(self) -> Dict[int, str]:
        """希望出勤。有給・研修などの「勤務しない」記号は含めない。"""
        return {
            day: mark
            for day, mark in sorted(self.entries.items())
            if day >= 1 and mark != WISH_OFF and mark not in ABSENCE_MARKS
        }

    def absence_days(self) -> Dict[int, str]:
        """有給・夏休・研修・健診。公休とは別に、その日を占める。"""
        return {
            day: mark
            for day, mark in sorted(self.entries.items())
            if day >= 1 and mark in ABSENCE_MARKS
        }


@dataclass
class DiscardedEntry:
    """月をまたぐ引き継ぎで、その月に存在しない日だったため捨てられた記入。"""

    name: str
    day: int
    mark: str


# --- 書き出し ------------------------------------------------------------------


def export_request_sheet(
    profiles: Sequence[StaffProfile],
    year: int,
    month: int,
    path: Path,
    existing: Sequence[StaffRequests] = (),
) -> List[DiscardedEntry]:
    """その月の希望入力シートを作る。existing を渡すと記入済みの内容を引き継ぐ。

    戻り値は、existing に含まれていたが月をまたいで日にちが存在しなかったため
    引き継げなかった記入の一覧。
    """
    try:
        days = month_days(year, month)
    except calendar.IllegalMonthError as error:
        raise RequestSheetError(f"月は1〜12で指定してください(指定値: {month})") from error

    day_numbers = {day.day for day in days}
    filled: Dict[str, Dict[int, str]] = {}
    discarded: List[DiscardedEntry] = []
    for request in existing:
        if request.staff_id in filled:
            raise RequestSheetError(
                f"希望の引き継ぎ元でスタッフIDが重複しています: {request.staff_id}"
                "(このままだと片方の記入が消えます。CSVの送信IDを確認してください)"
            )
        filled[request.staff_id] = request.entries
        for day, mark in request.entries.items():
            # CARRY_OVER_DAY(0日)は前月末の引き継ぎ欄で、その月の日にちではない。
            # 捨てられたわけではないので discarded に入れない。
            if day == CARRY_OVER_DAY:
                continue
            if day not in day_numbers:
                discarded.append(DiscardedEntry(name=request.name, day=day, mark=mark))

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME

    _write_intro(sheet, year, month, len(days))
    _write_header(sheet, days)

    # 休職中の人は、その月の希望を聞く必要がないので行を出さない
    active = [p for p in profiles if not p.is_on_leave(year, month)]
    for offset, profile in enumerate(active):
        row = _HEADER_ROW + 1 + offset
        _write_staff_row(sheet, row, profile, days, filled.get(profile.staff_id, {}))

    _add_validation(sheet, len(active), len(days))
    sheet.column_dimensions[get_column_letter(_ID_COLUMN)].hidden = True
    sheet.column_dimensions[get_column_letter(_NAME_COLUMN)].width = 12
    sheet.column_dimensions[get_column_letter(_CARRY_COLUMN)].width = 8
    sheet.freeze_panes = sheet.cell(row=_HEADER_ROW + 1, column=_FIRST_DAY_COLUMN)

    _write_legend(workbook)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return discarded


def _write_intro(sheet, year: int, month: int, day_count: int) -> None:
    sheet.cell(
        row=1, column=1, value=f"{month_label(year, month)} 希望休・希望出勤"
    ).font = Font(bold=True, size=14)
    sheet.cell(
        row=2,
        column=1,
        value=(
            "各セルのプルダウンから選んでください。"
            "休みたい日は「公」、この勤務に入りたい日はその記号を選びます。"
            "希望が無い日は空欄のままで結構です。"
            "いちばん左の「前月末」には、前月の最終日に○・△・◉だった方だけ入れてください。"
        ),
    )


def _write_header(sheet, days: Sequence[Day]) -> None:
    for label, column in (
        ("ID", _ID_COLUMN),
        ("スタッフ", _NAME_COLUMN),
        ("前月末", _CARRY_COLUMN),
    ):
        cell = sheet.cell(row=_HEADER_ROW, column=column, value=label)
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.fill = _HEAD_FILL
        cell.border = _BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for index, day in enumerate(days):
        column = _FIRST_DAY_COLUMN + index
        cell = sheet.cell(row=_HEADER_ROW, column=column, value=f"{day.day}\n{day.weekday_name}")
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = _HEAD_FILL
        cell.border = _BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(column)].width = 5


def _write_staff_row(
    sheet, row: int, profile: StaffProfile, days: Sequence[Day], entries: Dict[int, str]
) -> None:
    id_cell = sheet.cell(row=row, column=_ID_COLUMN, value=profile.staff_id)
    id_cell.border = _BORDER

    name_cell = sheet.cell(row=row, column=_NAME_COLUMN, value=profile.name)
    name_cell.fill = _NAME_FILL
    name_cell.border = _BORDER
    name_cell.alignment = Alignment(vertical="center")

    carry = sheet.cell(row=row, column=_CARRY_COLUMN)
    carry.value = entries.get(CARRY_OVER_DAY) or None
    carry.fill = _NAME_FILL
    carry.border = _BORDER
    carry.alignment = Alignment(horizontal="center", vertical="center")
    carry.number_format = "@"

    for index, day in enumerate(days):
        cell = sheet.cell(row=row, column=_FIRST_DAY_COLUMN + index)
        cell.value = entries.get(day.day) or None
        cell.border = _BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")
        # 文字列として扱わせる。Excelは「9-16」のような値を日付に変換してしまい、
        # 読み戻したときに希望が消える。
        cell.number_format = "@"
        if day.is_weekend:
            cell.fill = _WEEKEND_FILL


def _add_validation(sheet, staff_count: int, day_count: int) -> None:
    if not staff_count or not day_count:
        return
    carry = DataValidation(
        type="list", formula1='"' + ",".join(CARRY_OVER_CHOICES) + '"', allow_blank=True
    )
    sheet.add_data_validation(carry)
    carry_letter = get_column_letter(_CARRY_COLUMN)
    carry.add(
        f"{carry_letter}{_HEADER_ROW + 1}:{carry_letter}{_HEADER_ROW + staff_count}"
    )
    validation = DataValidation(
        type="list", formula1='"' + ",".join(WISH_CHOICES) + '"', allow_blank=True
    )
    sheet.add_data_validation(validation)
    first = get_column_letter(_FIRST_DAY_COLUMN)
    last = get_column_letter(_FIRST_DAY_COLUMN + day_count - 1)
    validation.add(f"{first}{_HEADER_ROW + 1}:{last}{_HEADER_ROW + staff_count}")


def _write_legend(workbook: Workbook) -> None:
    legend = workbook.create_sheet(LEGEND_SHEET)
    legend.column_dimensions["A"].width = 10
    legend.column_dimensions["B"].width = 60
    rows = [
        ("記号", "意味"),
        (OFF, "この日は休みたい（希望休）"),
        (
            "前月末",
            "前月の最終日の記号。○(夜勤)・△(明け)・◉(深夜)だけ入れます。"
            "前月から続く勤務を引き継ぐために使います",
        ),
    ]
    rows += [
        (code, f"この日は {times[0]}-{times[1]} の勤務に入りたい")
        for code, times in DAY_SHIFTS.items()
    ]
    rows += [
        (code, f"この日は {code} の時短で入りたい(勤務時間が決まっている方のみ)")
        for code in HOUR_CHOICES
    ]
    rows += [
        (NIGHT_IN, "この日は夜勤に入りたい（翌日は明け、その次は公休になります）"),
        (LATE_NIGHT_IN, "この日は深夜に入りたい（翌日は公休になります）"),
        ("有", "有給休暇（公休とは別に数えます）"),
        ("夏", "夏季休暇（公休とは別に数えます）"),
        ("研", "研修"),
        ("健", "健康診断"),
        ("(空欄)", "希望なし。こちらで組みます"),
    ]
    for index, (mark, meaning) in enumerate(rows, 1):
        cell = legend.cell(row=index, column=1, value=mark)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
        legend.cell(row=index, column=2, value=meaning)


# --- 読み戻し ------------------------------------------------------------------


def import_request_sheet(path: Path) -> Tuple[List[StaffRequests], List[str]]:
    """記入済みシートを読む。(スタッフごとの希望, 警告) を返す。"""
    workbook = load_workbook(path, data_only=True)
    if SHEET_NAME not in workbook.sheetnames:
        raise RequestSheetError(
            f"「{SHEET_NAME}」シートが見つかりません。シート名を変えないでください。"
        )
    sheet = workbook[SHEET_NAME]

    warnings: List[str] = []
    day_columns = _read_day_columns(sheet, warnings)
    if not day_columns:
        raise RequestSheetError(
            "日付の見出しが読み取れません。3行目の日付を消さないでください。"
        )

    requests: List[StaffRequests] = []
    allowed = set(WISH_CHOICES)

    for row in range(_HEADER_ROW + 1, sheet.max_row + 1):
        staff_id = _text(sheet.cell(row=row, column=_ID_COLUMN).value)
        name = _text(sheet.cell(row=row, column=_NAME_COLUMN).value)
        if not staff_id and not name:
            continue
        if not staff_id:
            warnings.append(f"{row}行目「{name}」はID列が空のため飛ばしました")
            continue

        entries: Dict[int, str] = {}

        carry = _text(sheet.cell(row=row, column=_CARRY_COLUMN).value)
        if carry:
            if carry in CARRY_OVER_CHOICES:
                entries[CARRY_OVER_DAY] = carry
            else:
                warnings.append(
                    f"{name} の前月末「{carry}」は"
                    f"{' / '.join(CARRY_OVER_CHOICES)} のどれかにしてください(無視しました)"
                )

        for day, column in day_columns.items():
            mark = _text(sheet.cell(row=row, column=column).value)
            if not mark:
                continue
            if mark not in allowed:
                warnings.append(
                    f"{name} の{day}日「{mark}」は選べる記号ではありません(無視しました)"
                )
                continue
            entries[day] = mark
        requests.append(StaffRequests(staff_id=staff_id, name=name, entries=entries))

    return requests, warnings


def _read_day_columns(sheet, warnings: List[str]) -> Dict[int, int]:
    """見出し行から {日にち: 列番号} を作る。列がずれても日付で対応付ける。"""
    columns: Dict[int, int] = {}
    for column in range(_FIRST_DAY_COLUMN, sheet.max_column + 1):
        value = _text(sheet.cell(row=_HEADER_ROW, column=column).value)
        head = value.split("\n")[0].strip()
        if head.isdigit():
            day = int(head)
            if day in columns:
                warnings.append(
                    f"{day}日の見出し列が複数あります"
                    f"(列{columns[day]}と列{column})。列{column}の内容を使いました"
                )
            columns[day] = column
    return columns


def _text(value) -> str:
    return str(value).strip() if value is not None else ""
