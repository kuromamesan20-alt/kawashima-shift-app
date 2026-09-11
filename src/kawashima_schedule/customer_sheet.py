"""お客様(施設)にお渡しする確認シートの書き出しと、回答の読み戻し。

Googleスプレッドシートで共有して、そのまま書き込んでいただく想定。
xlsx でお渡しし、記入後は xlsx でも CSV でも受け取れるようにする。
"""

from __future__ import annotations

import csv
from pathlib import Path
from typing import Dict, List, Sequence, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .questions import Question

SHEET_NAME = "ご確認のお願い"

_HEADERS = [
    "No",
    "スタッフ",
    "いただいたご記入",
    "ご確認いただきたいこと",
    "ご回答",
    "補足・ご自由にご記入ください",
    "整理番号",
]
_WIDTHS = [5, 12, 34, 52, 26, 34, 22]

# 回答欄(E列)と補足欄(F列)だけ記入していただく(書き出し時の並び。読み込み時は
# 見出し名で列位置を探すので、お客様が列を挿入・削除しても壊れない)
_ANSWER_COLUMN = 5
_NOTE_COLUMN = 6
_ID_COLUMN = 7

_ANSWER_HEADER = "ご回答"
_NOTE_HEADER = "補足・ご自由にご記入ください"
_ID_HEADER = "整理番号"

_HEAD_FILL = PatternFill("solid", fgColor="2F5D8C")
_LOCKED_FILL = PatternFill("solid", fgColor="F2F4F7")
_ANSWER_FILL = PatternFill("solid", fgColor="FFFDF5")
_THIN = Side(style="thin", color="D9DEE5")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


class AnswerSheetError(ValueError):
    """回答シートの形式が想定と違う場合に送出する。"""


def export_question_sheet(
    questions: Sequence[Question], path: Path, case_name: str = ""
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME

    _write_intro(sheet, case_name, len(questions))
    header_row = 5

    for index, header in enumerate(_HEADERS, 1):
        cell = sheet.cell(row=header_row, column=index, value=header)
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.fill = _HEAD_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = _BORDER
        sheet.column_dimensions[get_column_letter(index)].width = _WIDTHS[index - 1]

    for offset, question in enumerate(questions):
        row = header_row + 1 + offset
        question_text = question.text
        if question.free_text_hint:
            question_text = f"{question_text}\n（{question.free_text_hint}）"
        values = [
            offset + 1,
            question.staff_name,
            question.original_text,
            question_text,
            "",
            "",
            question.question_id,
        ]
        for index, value in enumerate(values, 1):
            cell = sheet.cell(row=row, column=index, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = _BORDER
            if index in (_ANSWER_COLUMN, _NOTE_COLUMN):
                cell.fill = _ANSWER_FILL
            else:
                cell.fill = _LOCKED_FILL

        if question.choices:
            _add_choice_validation(sheet, row, question.choices)
        else:
            sheet.cell(row=row, column=_ANSWER_COLUMN, value="").alignment = Alignment(
                vertical="top", wrap_text=True
            )

    # 整理番号の列は、こちらが回答を元に戻すためのもの。触らないでいただく。
    sheet.column_dimensions[get_column_letter(_ID_COLUMN)].hidden = True
    sheet.freeze_panes = f"A{header_row + 1}"
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def _write_intro(sheet, case_name: str, count: int) -> None:
    title = f"勤務表作成にあたってのご確認（{case_name}）" if case_name else "勤務表作成にあたってのご確認"
    sheet.cell(row=1, column=1, value=title).font = Font(bold=True, size=14)
    sheet.cell(
        row=2,
        column=1,
        value=(
            f"いただいたスタッフ情報のうち、{count}件について確認させてください。"
            "黄色い「ご回答」欄にご記入をお願いします。"
        ),
    )
    sheet.cell(
        row=3,
        column=1,
        value=(
            "「ご回答」欄はプルダウンから選べます。当てはまるものが無い場合や"
            "補足がある場合は、右の「補足」欄にご自由にお書きください。"
        ),
    )


def _add_choice_validation(sheet, row: int, choices: Sequence[str]) -> None:
    # Excelの入力規則は文字列長に上限があるため、長すぎる選択肢は素の記入欄にする
    formula = ",".join(choice.replace(",", "、") for choice in choices)
    if len(formula) > 250:
        return
    validation = DataValidation(type="list", formula1=f'"{formula}"', allow_blank=True)
    sheet.add_data_validation(validation)
    letter = get_column_letter(_ANSWER_COLUMN)
    validation.add(f"{letter}{row}")


# --- 回答の読み戻し --------------------------------------------------------------


def import_answer_sheet(path: Path) -> Dict[str, Dict[str, str]]:
    """記入済みの確認シートを読む。{整理番号: {"choice":…, "note":…}} を返す。"""
    if path.suffix.lower() == ".csv":
        rows = _read_csv(path)
    else:
        rows = _read_xlsx(path)

    answers: Dict[str, Dict[str, str]] = {}
    found_any_row = False
    for question_id, choice, note in rows:
        if not question_id:
            continue
        found_any_row = True
        if not choice and not note:
            continue  # 未記入の質問は「回答なし」として扱う
        answers[question_id] = {"choice": choice, "note": note}

    if not found_any_row:
        raise AnswerSheetError(
            "質問が1件も読み取れませんでした。"
            "「整理番号」の列が残っているシートをそのままお使いください。"
        )
    return answers


def _read_xlsx(path: Path) -> List[Tuple[str, str, str]]:
    workbook = load_workbook(path, data_only=True)
    sheet = workbook[SHEET_NAME] if SHEET_NAME in workbook.sheetnames else workbook.active
    all_rows = [[c.value for c in sheet[r]] for r in range(1, min(sheet.max_row, 20) + 1)]
    header_row = _find_header_row(all_rows)
    columns = _find_columns(all_rows[header_row - 1])
    return [
        (
            _text(sheet.cell(row=row, column=columns["id"]).value),
            _text(sheet.cell(row=row, column=columns["answer"]).value),
            _text(sheet.cell(row=row, column=columns["note"]).value),
        )
        for row in range(header_row + 1, sheet.max_row + 1)
    ]


def _read_csv(path: Path) -> List[Tuple[str, str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.reader(handle))
    # _find_header_row は1始まりの行番号を返す。リストは0始まりなので、
    # 見出しの次の行は rows[header_row] から始まる。
    header_row = _find_header_row(rows)
    columns = _find_columns(rows[header_row - 1])
    max_column = max(columns.values())
    result = []
    for row in rows[header_row:]:
        padded = row + [""] * (max_column - len(row))
        result.append(
            (
                _text(padded[columns["id"] - 1]),
                _text(padded[columns["answer"] - 1]),
                _text(padded[columns["note"] - 1]),
            )
        )
    return result


def _find_header_row(rows: Sequence[Sequence]) -> int:
    """見出し行(1始まり)を探す。案内文が何行あってもよいように。"""
    for index, row in enumerate(rows, 1):
        values = [_text(value) for value in row]
        if _ID_HEADER in values and _ANSWER_HEADER in values:
            return index
    raise AnswerSheetError(
        f"見出し行が見つかりません。「{_ID_HEADER}」「{_ANSWER_HEADER}」の列があるシートをお使いください。"
    )


def _find_columns(header_row: Sequence) -> Dict[str, int]:
    """見出し行から、必要な列の位置(1始まり)を名前で探す。

    お客様が列を挿入・削除・並べ替えても、正しい列を読めるようにするため。
    """
    values = [_text(value) for value in header_row]
    required = {
        "id": _ID_HEADER,
        "answer": _ANSWER_HEADER,
        "note": _NOTE_HEADER,
    }
    columns: Dict[str, int] = {}
    missing: List[str] = []
    for key, header_name in required.items():
        if header_name in values:
            columns[key] = values.index(header_name) + 1
        else:
            missing.append(header_name)
    if missing:
        raise AnswerSheetError(
            "次の列が見つかりませんでした。列の名前や見出し行を変更・削除していないか"
            f"ご確認ください: {'、'.join(missing)}"
        )
    return columns


def _text(value) -> str:
    return str(value).strip() if value is not None else ""
