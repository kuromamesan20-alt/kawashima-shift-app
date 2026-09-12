"""勤務表をExcelに書き出す。実物(勤務計画表_2026年07月_2病棟全体)と同じ体裁にする。

  A列  ユニット名(区画の先頭行にだけ入れる)
  B列  職種・肩書き
  C列  勤務表の番号
  D列  日付/曜日/行事 の見出し
  E〜  1日〜末日
  右端 日勤 / 夜勤 / 深夜 / 公休 / 残業 / 有休 / 夏正 / 出研 の集計

1人につき2行。1行目が勤務欄、2行目は中抜けの2コマ目と「せ」(責任者)を書く欄。
最下部に日別人数と、番号付き6種がそろっているかの判定を置く。
"""

from __future__ import annotations

from pathlib import Path
from typing import Dict, List, Sequence

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .calendar_utils import month_label
from .models import StaffProfile
from .scheduler import ScheduleResult
from .shifts import (
    DAY_SHIFTS,
    HEALTH_CHECK,
    LATE_NIGHT_IN,
    NIGHT_AFTER,
    NIGHT_IN,
    OFF,
    PAID_LEAVE,
    REQUIRED_DAY_SHIFTS,
    SUMMER_LEAVE,
    TRAINING,
)

SHEET_NAME = "勤務計画表"
RESPONSIBLE = "せ"

UNIT_COLUMN, ROLE_COLUMN, NUMBER_COLUMN, LABEL_COLUMN = 1, 2, 3, 4
FIRST_DAY_COLUMN = 5
HEADER_ROW, WEEKDAY_ROW, EVENT_ROW = 3, 4, 5
FIRST_STAFF_ROW = 6

SUMMARY = ("日勤", "夜勤", "深夜", "公休", "残業", "有休", "夏正", "出研")
# 集計の中身は数字ではなく Excel の計算式で入れる。
# この勤務表はたたき台で、受け取った側が手直しする前提のため、
# セルを書き換えたら右端と下の集計もその場で変わる必要がある。
#
# 「残業」は実物の勤務表にある手書き欄。記号からは決められないので空欄のまま。
# 「日勤」は番号付きの勤務だけを数える。時間を直接書く方のセル(「9-16時」など)は
# 数えない(施設の担当者に確認済み)。
SUMMARY_MARKS: Dict[str, tuple] = {
    "日勤": tuple(DAY_SHIFTS),
    "夜勤": (NIGHT_IN,),
    "深夜": (LATE_NIGHT_IN,),
    "公休": (OFF,),
    "残業": (),
    "有休": (PAID_LEAVE,),
    "夏正": (SUMMER_LEAVE,),
    "出研": (TRAINING, HEALTH_CHECK),
}
UNIT_ORDER = ("ばら", "さくら", "ゆり", "すみれ")

_HEAD_FILL = PatternFill("solid", fgColor="2F5D8C")
_WEEKEND_FILL = PatternFill("solid", fgColor="FDECEC")
_UNIT_FILL = PatternFill("solid", fgColor="EDF0F4")
_NIGHT_FILL = PatternFill("solid", fgColor="E4D7F5")
_OFF_FILL = PatternFill("solid", fgColor="F2F2F2")
_NG_FILL = PatternFill("solid", fgColor="FBE3E3")
_THIN = Side(style="thin", color="BFC7D1")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

# 記号ごとの色。目で追いやすくするためだけのもの。
_MARK_FILL = {
    NIGHT_IN: _NIGHT_FILL,
    NIGHT_AFTER: PatternFill("solid", fgColor="EFE8FA"),
    LATE_NIGHT_IN: PatternFill("solid", fgColor="D9D2E9"),
    OFF: _OFF_FILL,
    PAID_LEAVE: PatternFill("solid", fgColor="FFF2CC"),
    SUMMER_LEAVE: PatternFill("solid", fgColor="FFF2CC"),
    TRAINING: PatternFill("solid", fgColor="E2EFDA"),
    HEALTH_CHECK: PatternFill("solid", fgColor="E2EFDA"),
}


def export_schedule(
    result: ScheduleResult,
    profiles: Sequence[StaffProfile],
    path: Path,
    events: Dict[int, str] = None,
) -> None:
    """組み上がった勤務表をExcelにする。events は {日にち: 行事名}。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    days = result.days
    events = events or {}

    _write_legend_row(sheet, len(days))
    _write_title(sheet, result)
    _write_header(sheet, days, events)

    ordered = _order_staff(profiles, result)
    row = FIRST_STAFF_ROW
    previous_unit = None
    for profile in ordered:
        unit = _unit_label(profile)
        _write_staff(sheet, row, profile, result, unit if unit != previous_unit else "")
        previous_unit = unit
        row += 2

    staff_rows = (FIRST_STAFF_ROW, max(FIRST_STAFF_ROW, row - 1))
    _write_daily_check(sheet, row + 1, days, staff_rows)
    _finish_layout(sheet, days, len(ordered))

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


# --- 見出し -------------------------------------------------------------------


def _write_legend_row(sheet, day_count: int) -> None:
    """1行目にシフトの時間を並べる(実物と同じ)。"""
    column = FIRST_DAY_COLUMN
    for code, (start, end) in DAY_SHIFTS.items():
        cell = sheet.cell(row=1, column=column, value=f"{code} {start}-{end}")
        cell.font = Font(size=9)
        column += 4


def _write_title(sheet, result: ScheduleResult) -> None:
    sheet.cell(row=2, column=1, value="2病棟").font = Font(bold=True, size=12)
    sheet.cell(row=2, column=FIRST_DAY_COLUMN, value=month_label(result.year, result.month))


def _write_header(sheet, days, events: Dict[int, str]) -> None:
    for column, label in ((ROLE_COLUMN, "職種"), (LABEL_COLUMN, "日付")):
        _head(sheet, HEADER_ROW, column, label)
    _head(sheet, WEEKDAY_ROW, LABEL_COLUMN, "曜日")
    _head(sheet, EVENT_ROW, LABEL_COLUMN, "行事")

    for index, day in enumerate(days):
        column = FIRST_DAY_COLUMN + index
        _head(sheet, HEADER_ROW, column, day.day)
        _head(sheet, WEEKDAY_ROW, column, day.weekday_name)
        event = sheet.cell(row=EVENT_ROW, column=column, value=events.get(day.day))
        event.alignment = Alignment(horizontal="center", wrap_text=True)
        event.border = _BORDER
        event.font = Font(size=8)
        if day.is_weekend:
            event.fill = _WEEKEND_FILL

    for offset, label in enumerate(SUMMARY):
        _head(sheet, HEADER_ROW, FIRST_DAY_COLUMN + len(days) + offset, label)


def _head(sheet, row: int, column: int, value) -> None:
    cell = sheet.cell(row=row, column=column, value=value)
    cell.font = Font(bold=True, color="FFFFFF", size=10)
    cell.fill = _HEAD_FILL
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = _BORDER


# --- スタッフ行 ----------------------------------------------------------------


def _order_staff(profiles, result) -> List[StaffProfile]:
    """師長 → ユニット順 → 介護補助。勤務表に載らない人は除く。"""
    listed = [p for p in profiles if p.staff_id in result.assignments]

    def key(profile):
        if profile.is_head_nurse:
            return (0, "")
        if profile.is_support_staff:
            return (5, profile.sheet_label)
        rank = UNIT_ORDER.index(profile.unit) + 1 if profile.unit in UNIT_ORDER else 6
        return (rank, _number_key(profile.sheet_label))

    return sorted(listed, key=key)


def _number_key(label: str):
    """看護(1,2,3…)を先に、介護(①②③…)を後に並べる。"""
    circles = "①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮⑯⑰⑱⑲⑳㉑㉒㉓㉔㉕㉖"
    if label and label[0] in circles:
        # 丸数字は isdigit() でも真になるので、先に判定する
        return (1, circles.index(label[0]))
    if label.isdigit():
        return (0, int(label))
    return (2, label)


def _unit_label(profile: StaffProfile) -> str:
    if profile.is_head_nurse:
        return ""
    if profile.is_support_staff:
        return "介護補助"
    return profile.unit


def _write_staff(sheet, row: int, profile, result, unit_label: str) -> None:
    if unit_label:
        cell = sheet.cell(row=row, column=UNIT_COLUMN, value=unit_label)
        cell.font = Font(bold=True, size=10)
        cell.fill = _UNIT_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")

    role = sheet.cell(row=row, column=ROLE_COLUMN, value=_role_label(profile))
    role.font = Font(size=9)
    role.alignment = Alignment(vertical="center", wrap_text=True)
    role.border = _BORDER

    number = sheet.cell(row=row, column=NUMBER_COLUMN, value=profile.sheet_label)
    number.alignment = Alignment(horizontal="center", vertical="center")
    number.border = _BORDER

    sheet.cell(row=row, column=LABEL_COLUMN, value=profile.name).border = _BORDER
    sheet.cell(row=row + 1, column=LABEL_COLUMN).border = _BORDER

    assignment = result.assignments.get(profile.staff_id, {})

    for index, day in enumerate(result.days):
        column = FIRST_DAY_COLUMN + index
        mark = assignment.get(day.day, "")
        main, second = _split_mark(mark)

        _write_mark(sheet, row, column, main, day.is_weekend)
        extra = second
        if result.responsible.get(day.day) == profile.staff_id:
            extra = f"{extra} {RESPONSIBLE}".strip() if extra else RESPONSIBLE
        _write_mark(sheet, row + 1, column, extra, day.is_weekend, small=True)

    days_range = _row_range(row, len(result.days))
    for offset, label in enumerate(SUMMARY):
        cell = sheet.cell(
            row=row, column=FIRST_DAY_COLUMN + len(result.days) + offset,
            value=_countif_formula(days_range, SUMMARY_MARKS[label]),
        )
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _BORDER


def _role_label(profile: StaffProfile) -> str:
    if profile.is_head_nurse:
        return "師長"
    return profile.role or ""


def _split_mark(mark: str):
    """中抜け(「09:00-12:00・19:00-20:00」)は2行に分ける。"""
    if mark and "・" in mark:
        first, _, rest = mark.partition("・")
        return first, rest
    return mark, ""


def _write_mark(sheet, row: int, column: int, value: str, weekend: bool, small=False) -> None:
    cell = sheet.cell(row=row, column=column, value=value or None)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = _BORDER
    cell.font = Font(size=8 if small else 10)
    if value in _MARK_FILL:
        cell.fill = _MARK_FILL[value]
    elif weekend:
        cell.fill = _WEEKEND_FILL


def _row_range(row: int, day_count: int) -> str:
    """その人の勤務欄(1行分)の範囲。例: E6:AI6"""
    first = get_column_letter(FIRST_DAY_COLUMN)
    last = get_column_letter(FIRST_DAY_COLUMN + day_count - 1)
    return f"{first}{row}:{last}{row}"


def _countif_formula(cell_range: str, marks) -> str:
    """その範囲に marks がいくつあるかを数える式。marks が空なら空欄。

    数字ではなく式で入れるのは、受け取った側が勤務表を手直ししたときに
    集計がその場で変わるようにするため。
    """
    if not marks:
        return None
    return "=" + "+".join(f'COUNTIF({cell_range},"{mark}")' for mark in marks)


# --- 最下部の集計 --------------------------------------------------------------


def _write_daily_check(sheet, row: int, days, staff_rows) -> None:
    """番号付き6種が毎日そろっているかを確かめる行。実物にも同じものがある。

    こちらも数字ではなく式で入れる。勤務表を手直ししたときに、
    その日がそろっているかどうかがその場で分かるようにするため。

    数える範囲はスタッフ欄ぜんぶ(1人2行の2行目も含む)。2行目には中抜けの
    時間と「せ」しか入らず、番号付きの記号とは一致しないので混ざらない。
    """
    label = sheet.cell(row=row, column=ROLE_COLUMN, value="日別チェック")
    label.font = Font(bold=True, size=9)

    first_staff_row, last_staff_row = staff_rows
    codes = REQUIRED_DAY_SHIFTS + (NIGHT_IN, LATE_NIGHT_IN)
    count_rows = {}

    for offset, code in enumerate(codes):
        line = row + 1 + offset
        count_rows[code] = line
        sheet.cell(row=line, column=LABEL_COLUMN, value=code).alignment = Alignment(
            horizontal="center"
        )
        for index, day in enumerate(days):
            letter = get_column_letter(FIRST_DAY_COLUMN + index)
            cell = sheet.cell(
                row=line,
                column=FIRST_DAY_COLUMN + index,
                value=f'=COUNTIF({letter}{first_staff_row}:{letter}{last_staff_row},"{code}")',
            )
            cell.alignment = Alignment(horizontal="center")
            cell.font = Font(size=8)
            cell.border = _BORDER

    judge = row + 1 + len(codes) + 1
    sheet.cell(row=judge, column=ROLE_COLUMN, value="6種そろっているか").font = Font(
        bold=True, size=9
    )
    for index, day in enumerate(days):
        letter = get_column_letter(FIRST_DAY_COLUMN + index)
        conditions = ",".join(
            f"{letter}{count_rows[code]}=1" for code in REQUIRED_DAY_SHIFTS
        )
        cell = sheet.cell(
            row=judge,
            column=FIRST_DAY_COLUMN + index,
            value=f'=IF(AND({conditions}),"○","×")',
        )
        cell.alignment = Alignment(horizontal="center")
        cell.border = _BORDER

    _mark_shortage(sheet, judge, len(days))


def _mark_shortage(sheet, judge_row: int, day_count: int) -> None:
    """そろっていない日に色を付ける。

    色も式で決める(条件付き書式)。固定の色だと、手直ししてそろった後も
    赤いままになってしまう。
    """
    if not day_count:
        return
    first = get_column_letter(FIRST_DAY_COLUMN)
    last = get_column_letter(FIRST_DAY_COLUMN + day_count - 1)
    sheet.conditional_formatting.add(
        f"{first}{judge_row}:{last}{judge_row}",
        CellIsRule(operator="equal", formula=['"×"'], fill=_NG_FILL),
    )


def _finish_layout(sheet, days, staff_count: int) -> None:
    sheet.column_dimensions["A"].width = 9
    sheet.column_dimensions["B"].width = 16
    sheet.column_dimensions["C"].width = 5
    sheet.column_dimensions["D"].width = 9
    for index in range(len(days)):
        sheet.column_dimensions[get_column_letter(FIRST_DAY_COLUMN + index)].width = 4.5
    for offset in range(len(SUMMARY)):
        sheet.column_dimensions[
            get_column_letter(FIRST_DAY_COLUMN + len(days) + offset)
        ].width = 5.5
    sheet.freeze_panes = sheet.cell(row=FIRST_STAFF_ROW, column=FIRST_DAY_COLUMN)
