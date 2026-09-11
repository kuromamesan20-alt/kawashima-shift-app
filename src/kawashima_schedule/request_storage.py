"""希望休・希望出勤の保存先。

本番(Streamlit Cloud)では Googleスプレッドシート、
設定が無いローカルでは data/staff_profiles/ のYAMLに保存する。
どちらでも同じ関数で読み書きできるようにしておく。
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence

from openpyxl.utils import get_column_letter

from .request_sheet import StaffRequests

SCOPES = [
    "https://www.googleapis.com/auth/spreadsheets",
    "https://www.googleapis.com/auth/drive",
]
DEFAULT_WORKSHEET = "requests"
# スタッフ情報を置くシート。1行1人で、条件は YAML のまま1セルに入れる。
# 項目が40近くあり今後も増えるため、列に展開せず丸ごと持たせる。
PROFILES_WORKSHEET = "staff_profiles"
PROFILE_HEADER = ["スタッフID", "名前", "勤務条件(YAML)"]
HEADER = ["年", "月", "スタッフID", "名前", "日", "記号"]


@dataclass(frozen=True)
class StorageStatus:
    """いま何処に保存しているかを画面に出すためのもの。"""

    provider: str
    is_cloud: bool
    message: str


class RequestStorage:
    """保存先の共通の形。"""

    def load(self, year: int, month: int) -> List[StaffRequests]:
        raise NotImplementedError

    def save(self, year: int, month: int, requests: Sequence[StaffRequests]) -> None:
        raise NotImplementedError

    def status(self) -> StorageStatus:
        raise NotImplementedError

    def load_profiles(self):
        """スタッフ情報を読む。置き場所が無ければ None を返す。"""
        return None

    def save_profiles(self, profiles) -> None:
        raise NotImplementedError


# --- ローカル(YAML) -------------------------------------------------------------


class LocalStorage(RequestStorage):
    """手元で動かすとき用。data/staff_profiles/YYYY-MM_希望.yaml に保存する。"""

    def __init__(self, directory: Path):
        self.directory = Path(directory)

    def _path(self, year: int, month: int) -> Path:
        return self.directory / f"{year}-{month:02d}_希望.yaml"

    def load(self, year: int, month: int) -> List[StaffRequests]:
        from .profile_store import load_requests

        path = self._path(year, month)
        return load_requests(path) if path.exists() else []

    def save(self, year: int, month: int, requests: Sequence[StaffRequests]) -> None:
        from .profile_store import save_requests

        save_requests(list(requests), self._path(year, month))

    def load_profiles(self):
        from .profile_store import load_profiles as _load

        path = self.directory / "case001.yaml"
        return _load(path) if path.exists() else None

    def save_profiles(self, profiles) -> None:
        from .profile_store import save_profiles as _save

        _save(list(profiles), self.directory / "case001.yaml")

    def status(self) -> StorageStatus:
        return StorageStatus(
            provider="この端末のファイル",
            is_cloud=False,
            message=f"{self.directory} に保存します(この端末の中だけ)",
        )


# --- Googleスプレッドシート -------------------------------------------------------


class GoogleSheetStorage(RequestStorage):
    """Streamlit Cloud 用。1行=1人1日の記号、という形で持つ。

    月をまたいでも同じシートに貯めるので、保存時はその月の行だけ入れ替える。
    """

    def __init__(self, sheet_id: str, credentials: Dict[str, Any], worksheet: str = DEFAULT_WORKSHEET):
        self.sheet_id = sheet_id
        self.credentials = credentials
        self.worksheet_name = worksheet

    def _book(self):
        import gspread
        from google.oauth2.service_account import Credentials

        creds = Credentials.from_service_account_info(self.credentials, scopes=SCOPES)
        return gspread.authorize(creds).open_by_key(self.sheet_id)

    def _sheet(self, name: str, header: List[str], create: bool = True):
        """名前でシートを取り、見出しを整える。

        create=False のときは、無ければ None を返す(まだ用意していない状態)。
        """
        book = self._book()
        try:
            sheet = book.worksheet(name)
        except Exception:
            if not create:
                return None
            sheet = book.add_worksheet(name, rows=2000, cols=len(header))
            sheet.update("A1", [header])
            return sheet
        if sheet.row_values(1) != header:
            sheet.update("A1", [header])
        return sheet

    def _worksheet(self):
        return self._sheet(self.worksheet_name, HEADER, create=True)

    def load(self, year: int, month: int) -> List[StaffRequests]:
        rows = self._worksheet().get_all_records()
        by_staff: Dict[str, StaffRequests] = {}
        for row in rows:
            if str(row.get("年")) != str(year) or str(row.get("月")) != str(month):
                continue
            staff_id = str(row.get("スタッフID", "")).strip()
            if not staff_id:
                continue
            entry = by_staff.setdefault(
                staff_id,
                StaffRequests(staff_id=staff_id, name=str(row.get("名前", "")).strip()),
            )
            try:
                day = int(row.get("日"))
            except (TypeError, ValueError):
                continue
            mark = str(row.get("記号", "")).strip()
            if mark:
                entry.entries[day] = mark
        return list(by_staff.values())

    def save(self, year: int, month: int, requests: Sequence[StaffRequests]) -> None:
        sheet = self._worksheet()
        existing = sheet.get_all_records()
        previous_total_rows = len(existing) + 1  # ヘッダー分を足す

        # その月以外の行は残し、対象月だけ入れ替える
        kept = [
            [r.get("年"), r.get("月"), r.get("スタッフID"), r.get("名前"), r.get("日"), r.get("記号")]
            for r in existing
            if not (str(r.get("年")) == str(year) and str(r.get("月")) == str(month))
        ]
        fresh = [
            [year, month, request.staff_id, request.name, day, mark]
            for request in requests
            for day, mark in sorted(request.entries.items())
            if mark
        ]
        rows = [HEADER] + kept + fresh

        # 「消してから書く」と、書き込み失敗時に過去分ごと全データを失うため、
        # 先に上書きし、はみ出した古い行だけ後から空にする順番にする。
        self._ensure_capacity(sheet, len(rows))
        last_row = len(rows)
        last_column = get_column_letter(len(HEADER))
        sheet.update(f"A1:{last_column}{last_row}", rows)

        if previous_total_rows > last_row:
            blank = [[""] * len(HEADER) for _ in range(previous_total_rows - last_row)]
            sheet.update(f"A{last_row + 1}:{last_column}{previous_total_rows}", blank)

    def _ensure_capacity(self, sheet, needed_rows: int) -> None:
        """行数上限に引っかからないよう、書き込み前に余裕を持たせて拡張する。

        34名×31日=最大1054行/月を毎月貯めていくと、既定の2000行では
        数ヶ月で上限に達してしまう。ここでまとめて先に拡張しておく。
        """
        buffer_rows = 3000
        target = needed_rows + buffer_rows
        if getattr(sheet, "row_count", 0) < target:
            sheet.resize(rows=target)

    def load_profiles(self):
        """スタッフ情報をスプレッドシートから読む。

        1行1人。勤務条件は YAML のまま1セルに入っているので、
        そのまま組み立て直す。シートが無ければ None(未設定)。
        """
        import yaml

        try:
            sheet = self._sheet(PROFILES_WORKSHEET, PROFILE_HEADER, create=False)
        except Exception:
            return None
        if sheet is None:
            return None

        from .profile_store import _from_dict

        profiles = []
        for row in sheet.get_all_records():
            body = str(row.get("勤務条件(YAML)", "")).strip()
            if not body:
                continue
            try:
                entry = yaml.safe_load(body)
            except yaml.YAMLError:
                continue
            if isinstance(entry, dict):
                profiles.append(_from_dict(entry))
        return profiles or None

    def save_profiles(self, profiles) -> None:
        """スタッフ情報をスプレッドシートに書き出す。"""
        import yaml
        from dataclasses import asdict

        from .profile_store import _to_dict

        sheet = self._sheet(PROFILES_WORKSHEET, PROFILE_HEADER, create=True)
        rows = [PROFILE_HEADER]
        for profile in profiles:
            body = yaml.safe_dump(
                _to_dict(profile), allow_unicode=True, sort_keys=False, width=10000
            )
            rows.append([profile.staff_id, profile.name, body])

        self._ensure_capacity(sheet, len(rows))
        last = get_column_letter(len(PROFILE_HEADER))
        sheet.update(f"A1:{last}{len(rows)}", rows)

    def status(self) -> StorageStatus:
        return StorageStatus(
            provider="Googleスプレッドシート",
            is_cloud=True,
            message="入力内容はスプレッドシートに保存されます",
        )


# --- 保存先の選択 ----------------------------------------------------------------


def get_storage(secrets: Any = None, local_dir: Path = None) -> RequestStorage:
    """Secrets が揃っていればスプレッドシート、無ければこの端末のファイル。"""
    config = _google_config(secrets)
    if config:
        return GoogleSheetStorage(**config)
    return LocalStorage(local_dir or Path("data/staff_profiles"))


def _google_config(secrets: Any) -> Optional[Dict[str, Any]]:
    """Secrets から設定を取り出す。

    Streamlit の st.secrets は、secrets.toml が無い環境だと真偽判定や .get で
    例外を投げる。手元で動かすときは設定が無いのが普通なので、
    例外は「設定なし」として扱い、ローカル保存に落とす。
    """
    if secrets is None:
        return None
    try:
        sheet_id = secrets.get("requests_sheet_id")
        account = secrets.get("gcp_service_account")
    except Exception:
        return None
    if not sheet_id or not account:
        return None
    return {
        "sheet_id": str(sheet_id),
        "credentials": dict(account),
        "worksheet": str(secrets.get("requests_worksheet_name") or DEFAULT_WORKSHEET),
    }
