"""お客様確認シートの書き出しと、回答の読み戻しのテスト。

大事なのは「回答を推測で反映しない」こと。
はっきり「はい」と答えていただいた項目だけ確認済みにし、
それ以外は記録するだけで、確認事項は残す。
"""

from __future__ import annotations

import csv
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).parent.parent / "src"))

from kawashima_schedule.customer_sheet import (  # noqa: E402
    SHEET_NAME,
    AnswerSheetError,
    export_question_sheet,
    import_answer_sheet,
)
from kawashima_schedule.models import PairConstraint, ReviewItem, StaffProfile  # noqa: E402
from kawashima_schedule.questions import YES, apply_answers, build_questions  # noqa: E402


def sample_profiles():
    return [
        StaffProfile(
            staff_id="id-1",
            name="スタッフさ",
            raw_conditions="夜勤スタッフEとWは同席不可。",
            pair_constraints=[
                PairConstraint(other_staff="スタッフE", kind="no_pair_night")
            ],
            review_items=[
                ReviewItem(
                    message="同席制約を確認してください",
                    code="pair",
                    payload={
                        "sentence": "夜勤スタッフEとWは同席不可",
                        "others": ["スタッフE", "スタッフW"],
                        "limit": None,
                    },
                ),
                ReviewItem(
                    message="未解釈の記述あり",
                    code="unparsed",
                    payload={"sentence": "介護リーダー"},
                ),
            ],
        )
    ]


def fill(path: Path, answers):
    """確認シートの回答欄を埋める。answers は {整理番号: (回答, 補足)}。"""
    workbook = load_workbook(path)
    sheet = workbook[SHEET_NAME]
    for row in range(1, sheet.max_row + 1):
        question_id = sheet.cell(row=row, column=7).value
        if question_id in answers:
            choice, note = answers[question_id]
            sheet.cell(row=row, column=5, value=choice)
            sheet.cell(row=row, column=6, value=note)
    workbook.save(path)


# --- 質問の組み立て -------------------------------------------------------------


def test_questions_are_written_for_the_facility_not_for_us():
    """お客様向けの文面になっていること(内部用語を出さない)。"""
    questions = build_questions(sample_profiles())
    assert len(questions) == 2

    pair_question = questions[0]
    assert "スタッフE" in pair_question.text
    assert "スタッフW" in pair_question.text
    assert pair_question.choices, "同席の確認は選択式にする"

    for question in questions:
        for internal_word in ("YAML", "CSV", "review", "profile", "パース"):
            assert internal_word not in question.text


def test_question_ids_are_unique():
    questions = build_questions(sample_profiles())
    assert len({q.question_id for q in questions}) == len(questions)


# --- 書き出し・読み戻し ---------------------------------------------------------


def test_round_trip_xlsx(tmp_path):
    path = tmp_path / "questions.xlsx"
    questions = build_questions(sample_profiles())
    export_question_sheet(questions, path, "案件001")

    fill(path, {questions[0].question_id: (YES, ""), questions[1].question_id: ("", "介護のリーダーという意味です")})
    answers = import_answer_sheet(path)

    assert answers[questions[0].question_id]["choice"] == YES
    assert "介護のリーダー" in answers[questions[1].question_id]["note"]


def test_round_trip_csv(tmp_path):
    """Googleスプレッドシートから CSV で書き出して戻す場合。"""
    path = tmp_path / "questions.csv"
    questions = build_questions(sample_profiles())
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["案内文の行"])
        writer.writerow(
            ["No", "スタッフ", "いただいたご記入", "ご確認いただきたいこと", "ご回答", "補足・ご自由にご記入ください", "整理番号"]
        )
        writer.writerow(["1", "スタッフさ", "原文", "質問", YES, "", questions[0].question_id])

    answers = import_answer_sheet(path)
    assert answers[questions[0].question_id]["choice"] == YES


def test_sheet_without_id_column_is_rejected(tmp_path):
    path = tmp_path / "broken.csv"
    path.write_text("No,スタッフ,ご回答\n1,スタッフさ,はい\n", encoding="utf-8")
    with pytest.raises(AnswerSheetError):
        import_answer_sheet(path)


def test_import_reads_columns_by_header_name_even_if_reordered(tmp_path):
    """お客様が列を挿入して並びが変わっても、見出し名で正しく読めること。"""
    path = tmp_path / "reordered.csv"
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["案内文の行"])
        # 「メモ」列が「ご回答」の前に1本挿入されている
        writer.writerow(
            [
                "No",
                "スタッフ",
                "いただいたご記入",
                "ご確認いただきたいこと",
                "メモ",
                "ご回答",
                "補足・ご自由にご記入ください",
                "整理番号",
            ]
        )
        writer.writerow(["1", "スタッフさ", "原文", "質問", "", YES, "", "id-1#abc"])

    answers = import_answer_sheet(path)
    assert answers["id-1#abc"]["choice"] == YES


def test_import_missing_id_column_is_rejected_with_specific_message(tmp_path):
    """「整理番号」列が削除されたシートは、その旨のエラーになること。"""
    path = tmp_path / "no_id_column.csv"
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["案内文の行"])
        writer.writerow(
            ["No", "スタッフ", "いただいたご記入", "ご確認いただきたいこと", "ご回答", "補足・ご自由にご記入ください"]
        )
        writer.writerow(["1", "スタッフさ", "原文", "質問", YES, ""])

    with pytest.raises(AnswerSheetError) as excinfo:
        import_answer_sheet(path)
    assert "整理番号" in str(excinfo.value)


def test_export_reordered_columns_are_still_read_correctly_xlsx(tmp_path):
    """xlsx側でも、列を1本挿入した状態で正しく読めること。"""
    path = tmp_path / "questions.xlsx"
    questions = build_questions(sample_profiles())
    export_question_sheet(questions, path, "案件001")

    workbook = load_workbook(path)
    sheet = workbook[SHEET_NAME]
    sheet.insert_cols(5)  # 「ご確認いただきたいこと」と「ご回答」の間に1列挿入
    sheet.cell(row=5, column=5, value="メモ")

    # 挿入後の「ご回答」「整理番号」の実際の列位置を見出し名から探して記入する
    header_row = 5
    headers = {sheet.cell(row=header_row, column=c).value: c for c in range(1, sheet.max_column + 1)}
    answer_col = headers["ご回答"]
    id_col = headers["整理番号"]
    for row in range(header_row + 1, sheet.max_row + 1):
        if sheet.cell(row=row, column=id_col).value == questions[0].question_id:
            sheet.cell(row=row, column=answer_col, value=YES)
    workbook.save(path)

    answers = import_answer_sheet(path)
    assert answers[questions[0].question_id]["choice"] == YES


# --- 回答の反映 -----------------------------------------------------------------


def test_yes_resolves_the_item():
    profiles = sample_profiles()
    questions = build_questions(profiles)
    apply_answers(profiles, {questions[0].question_id: {"choice": YES, "note": ""}})

    assert len(profiles[0].review_items) == 1
    assert profiles[0].review_items[0].code == "unparsed", "残るのは未解釈の方"


def test_multiple_yes_answers_remove_the_right_items():
    """複数「はい」でも番号がずれて別の項目を消さないこと。"""
    profiles = [
        StaffProfile(
            staff_id="id-1",
            name="スタッフさ",
            review_items=[
                ReviewItem(
                    message="早出の解釈",
                    code="early_late_column",
                    payload={"label": "早出", "removed": ["日①", "日②"]},
                ),
                ReviewItem(
                    message="未解釈の記述あり",
                    code="unparsed",
                    payload={"sentence": "介護リーダー"},
                ),
                ReviewItem(
                    message="遅出の解釈",
                    code="early_late_column",
                    payload={"label": "遅出", "removed": ["日⑦", "日⑨"]},
                ),
            ],
        )
    ]
    questions = build_questions(profiles)
    apply_answers(
        profiles,
        {
            questions[0].question_id: {"choice": YES, "note": ""},
            questions[2].question_id: {"choice": YES, "note": ""},
        },
    )
    assert [item.code for item in profiles[0].review_items] == ["unparsed"], (
        "「はい」と答えた1番目と3番目だけが消え、間の2番目が残ること"
    )


def test_yes_typed_into_a_free_text_answer_does_not_resolve():
    """自由記入の欄に「はい」と書かれても、勝手に確認済みにしないこと。"""
    profiles = sample_profiles()
    questions = build_questions(profiles)
    assert not questions[1].choices, "この質問は自由記入のはず"

    apply_answers(profiles, {questions[1].question_id: {"choice": YES, "note": ""}})
    assert len(profiles[0].review_items) == 2


def test_free_text_answer_is_recorded_but_not_applied():
    """自由記入の回答は推測で反映せず、記録して確認事項は残す。"""
    profiles = sample_profiles()
    questions = build_questions(profiles)
    apply_answers(
        profiles,
        {questions[1].question_id: {"choice": "", "note": "夜勤は月3回までです"}},
    )

    assert profiles[0].customer_answers[questions[1].question_id] == "夜勤は月3回までです"
    assert len(profiles[0].review_items) == 2, "回答をもらっても確認事項は残す"


def test_yes_with_a_note_is_not_auto_resolved():
    """「はい」でも補足が書かれていたら、内容を見るまで確認済みにしない。"""
    profiles = sample_profiles()
    questions = build_questions(profiles)
    apply_answers(
        profiles,
        {questions[0].question_id: {"choice": YES, "note": "ただしWは月1回まで可"}},
    )
    assert len(profiles[0].review_items) == 2


def test_unknown_staff_id_is_reported():
    profiles = sample_profiles()
    summary = apply_answers(profiles, {"missing-id#1": {"choice": YES, "note": ""}})
    assert any("見つかりません" in line for line in summary)


# --- free_text_hint の表示 --------------------------------------------------------


def test_free_text_hint_is_visible_in_the_exported_sheet(tmp_path):
    """記入案内(free_text_hint)が、実際にお客様の見える場所に出ていること。"""
    path = tmp_path / "questions.xlsx"
    questions = build_questions(sample_profiles())
    assert any(q.free_text_hint for q in questions), "テストの前提として hint がある質問が必要"

    export_question_sheet(questions, path, "案件001")
    workbook = load_workbook(path)
    sheet = workbook[SHEET_NAME]

    header_row = 5
    for offset, question in enumerate(questions):
        if not question.free_text_hint:
            continue
        row = header_row + 1 + offset
        cell_text = sheet.cell(row=row, column=4).value or ""
        assert question.free_text_hint in cell_text
