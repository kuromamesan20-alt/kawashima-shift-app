"""希望休・希望出勤シートの書き出し・読み戻しのテスト。"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).parent.parent / "src"))
sys.path.insert(0, str(Path(__file__).parent.parent))

from kawashima_schedule.calendar_utils import month_days  # noqa: E402
from kawashima_schedule.models import StaffProfile  # noqa: E402
from kawashima_schedule.profile_store import (  # noqa: E402
    load_requests,
    save_profiles,
    save_requests,
)
from kawashima_schedule.request_sheet import (  # noqa: E402
    SHEET_NAME,
    RequestSheetError,
    StaffRequests,
    export_request_sheet,
    import_request_sheet,
)

import cli  # noqa: E402

PROFILES = [
    StaffProfile(staff_id="id-1", name="スタッフA"),
    StaffProfile(staff_id="id-2", name="スタッフB"),
]


def write_cell(path: Path, staff_row: int, day: int, value: str) -> None:
    workbook = load_workbook(path)
    sheet = workbook[SHEET_NAME]
    for column in range(3, sheet.max_column + 1):
        head = str(sheet.cell(row=3, column=column).value or "").split("\n")[0]
        if head == str(day):
            sheet.cell(row=staff_row, column=column, value=value)
            workbook.save(path)
            return
    raise AssertionError(f"{day}日の列がありません")


# --- 日付 ---------------------------------------------------------------------


def test_month_days_covers_the_whole_month():
    days = month_days(2026, 10)
    assert len(days) == 31
    assert days[0].day == 1 and days[0].weekday_name == "木"
    assert days[-1].day == 31


def test_february_in_a_leap_year():
    assert len(month_days(2028, 2)) == 29


# --- 書き出し・読み戻し ---------------------------------------------------------


def test_round_trip(tmp_path):
    path = tmp_path / "requests.xlsx"
    export_request_sheet(PROFILES, 2026, 10, path)

    write_cell(path, 4, 3, "公")  # スタッフA: 3日は休みたい
    write_cell(path, 4, 12, "○")  # スタッフA: 12日は夜勤希望
    write_cell(path, 5, 20, "日③")  # スタッフB: 20日は日③希望

    requests, warnings = import_request_sheet(path)
    assert warnings == []

    by_name = {r.name: r for r in requests}
    assert by_name["スタッフA"].wish_off_days() == [3]
    assert by_name["スタッフA"].wish_work_days() == {12: "○"}
    assert by_name["スタッフB"].wish_work_days() == {20: "日③"}
    assert by_name["スタッフB"].wish_off_days() == []


def test_blank_sheet_reads_as_no_requests(tmp_path):
    path = tmp_path / "requests.xlsx"
    export_request_sheet(PROFILES, 2026, 10, path)

    requests, warnings = import_request_sheet(path)
    assert warnings == []
    assert len(requests) == 2
    assert all(not r.entries for r in requests)


def test_unknown_mark_is_warned_not_guessed(tmp_path):
    path = tmp_path / "requests.xlsx"
    export_request_sheet(PROFILES, 2026, 10, path)
    write_cell(path, 4, 5, "やすみ")

    requests, warnings = import_request_sheet(path)
    assert any("やすみ" in w for w in warnings)
    assert requests[0].entries == {}, "読めない記号は取り込まない"


def test_columns_are_matched_by_date_not_position(tmp_path):
    """列が1本増えても、日付の見出しで対応付けられること。"""
    path = tmp_path / "requests.xlsx"
    export_request_sheet(PROFILES, 2026, 10, path)
    write_cell(path, 4, 7, "公")

    workbook = load_workbook(path)
    sheet = workbook[SHEET_NAME]
    sheet.insert_cols(3)  # 日付列の手前に1本挿入
    workbook.save(path)

    requests, _ = import_request_sheet(path)
    assert requests[0].wish_off_days() == [7]


def test_sheet_without_the_expected_name_is_rejected(tmp_path):
    path = tmp_path / "requests.xlsx"
    export_request_sheet(PROFILES, 2026, 10, path)

    workbook = load_workbook(path)
    workbook[SHEET_NAME].title = "べつの名前"
    workbook.save(path)

    with pytest.raises(RequestSheetError):
        import_request_sheet(path)


def test_carry_over_keeps_what_was_already_filled(tmp_path):
    """翌月分を作るときに、記入済みの内容を引き継げること。"""
    first = tmp_path / "first.xlsx"
    export_request_sheet(PROFILES, 2026, 10, first)
    write_cell(first, 4, 3, "公")

    existing, _ = import_request_sheet(first)
    second = tmp_path / "second.xlsx"
    export_request_sheet(PROFILES, 2026, 10, second, existing)

    requests, _ = import_request_sheet(second)
    assert requests[0].wish_off_days() == [3]


def test_shorter_month_drops_days_that_do_not_exist(tmp_path):
    """31日の希望を30日しかない月に持ち越しても、無い日は消えること。"""
    existing = [StaffRequests(staff_id="id-1", name="スタッフA", entries={31: "公"})]
    path = tmp_path / "nov.xlsx"
    export_request_sheet(PROFILES, 2026, 11, path)  # 11月は30日まで
    export_request_sheet(PROFILES, 2026, 11, path, existing)

    requests, _ = import_request_sheet(path)
    assert requests[0].entries == {}


def test_shorter_month_reports_discarded_entries(tmp_path):
    """月をまたぐ引き継ぎで捨てられた記入が戻り値でわかること。"""
    existing = [StaffRequests(staff_id="id-1", name="スタッフA", entries={31: "公"})]
    path = tmp_path / "nov.xlsx"

    discarded = export_request_sheet(PROFILES, 2026, 11, path, existing)  # 11月は30日まで

    assert len(discarded) == 1
    assert discarded[0].name == "スタッフA"
    assert discarded[0].day == 31
    assert discarded[0].mark == "公"


def test_export_request_sheet_rejects_duplicate_staff_id_in_existing(tmp_path):
    """existing の staff_id が重複していると記入が黙って消えるので、エラーにする。"""
    existing = [
        StaffRequests(staff_id="id-1", name="スタッフA", entries={3: "公"}),
        StaffRequests(staff_id="id-1", name="別のスタッフA", entries={4: "○"}),
    ]
    path = tmp_path / "out.xlsx"

    with pytest.raises(RequestSheetError, match="id-1"):
        export_request_sheet(PROFILES, 2026, 10, path, existing)


def test_export_request_sheet_rejects_out_of_range_month(tmp_path):
    path = tmp_path / "out.xlsx"
    with pytest.raises(RequestSheetError, match="1〜12"):
        export_request_sheet(PROFILES, 2026, 13, path)


def test_duplicate_day_heading_is_warned(tmp_path):
    """日付の見出し列が2本あると、黙って後勝ちにせず警告すること。"""
    path = tmp_path / "requests.xlsx"
    export_request_sheet(PROFILES, 2026, 10, path)

    workbook = load_workbook(path)
    sheet = workbook[SHEET_NAME]
    original_first_day_value = sheet.cell(row=3, column=3).value
    sheet.insert_cols(3)
    sheet.cell(row=3, column=3, value=original_first_day_value)  # 1日の見出しを複製
    workbook.save(path)

    _, warnings = import_request_sheet(path)
    assert any("1日" in w for w in warnings)


# --- cli.cmd_export_requests / cmd_import_requests --------------------------


def test_cmd_export_requests_refuses_to_overwrite_filled_sheet(tmp_path, capsys):
    profiles_path = tmp_path / "profiles.yaml"
    save_profiles(PROFILES, profiles_path)

    out_path = tmp_path / "out.xlsx"
    export_request_sheet(PROFILES, 2026, 10, out_path)
    write_cell(out_path, 4, 3, "公")  # 記入済みにする

    args = argparse.Namespace(
        profiles=str(profiles_path),
        year=2026,
        month=10,
        out=str(out_path),
        carry_over=None,
        force=False,
    )
    result = cli.cmd_export_requests(args)

    assert result == 1
    err = capsys.readouterr().err
    assert "--force" in err
    assert "--carry-over" in err


def test_cmd_export_requests_force_allows_overwrite(tmp_path):
    profiles_path = tmp_path / "profiles.yaml"
    save_profiles(PROFILES, profiles_path)

    out_path = tmp_path / "out.xlsx"
    export_request_sheet(PROFILES, 2026, 10, out_path)
    write_cell(out_path, 4, 3, "公")  # 記入済みにする

    args = argparse.Namespace(
        profiles=str(profiles_path),
        year=2026,
        month=10,
        out=str(out_path),
        carry_over=None,
        force=True,
    )
    result = cli.cmd_export_requests(args)

    assert result == 0
    requests, _ = import_request_sheet(out_path)
    assert all(not r.entries for r in requests), "強制上書きなので記入は消える"


def test_cmd_export_requests_allows_overwrite_when_target_is_empty(tmp_path):
    profiles_path = tmp_path / "profiles.yaml"
    save_profiles(PROFILES, profiles_path)

    out_path = tmp_path / "out.xlsx"
    export_request_sheet(PROFILES, 2026, 10, out_path)  # 記入なし

    args = argparse.Namespace(
        profiles=str(profiles_path),
        year=2026,
        month=10,
        out=str(out_path),
        carry_over=None,
        force=False,
    )
    result = cli.cmd_export_requests(args)

    assert result == 0


def test_cmd_export_requests_rejects_invalid_month(tmp_path, capsys):
    profiles_path = tmp_path / "profiles.yaml"
    save_profiles(PROFILES, profiles_path)

    args = argparse.Namespace(
        profiles=str(profiles_path),
        year=2026,
        month=13,
        out=str(tmp_path / "out.xlsx"),
        carry_over=None,
        force=False,
    )
    result = cli.cmd_export_requests(args)

    assert result == 1
    assert "1〜12" in capsys.readouterr().err


def test_cmd_import_requests_warns_about_row_missing_from_sheet(tmp_path, capsys):
    profiles = list(PROFILES) + [StaffProfile(staff_id="id-3", name="スタッフC")]
    profiles_path = tmp_path / "profiles.yaml"
    save_profiles(profiles, profiles_path)

    sheet_path = tmp_path / "requests.xlsx"
    export_request_sheet(PROFILES, 2026, 10, sheet_path)  # id-3 の行は無い

    args = argparse.Namespace(
        requests=str(sheet_path),
        out=str(tmp_path / "out.yaml"),
        profiles=str(profiles_path),
    )
    result = cli.cmd_import_requests(args)

    assert result == 0
    out = capsys.readouterr().out
    assert "スタッフC" in out
    assert "シートにありません" in out


# --- YAML保存 ------------------------------------------------------------------


def test_requests_round_trip_through_yaml(tmp_path):
    original = [
        StaffRequests(staff_id="id-1", name="スタッフA", entries={3: "公", 12: "○"})
    ]
    path = tmp_path / "requests.yaml"
    save_requests(original, path)

    loaded = load_requests(path)
    assert loaded[0].staff_id == "id-1"
    assert loaded[0].entries == {3: "公", 12: "○"}


# --- 有給・夏休・研修・健診 -------------------------------------------------------


def test_absence_marks_are_separate_from_wishes(tmp_path):
    """有給などは希望出勤にも希望休にも数えず、別に取り出せること。"""
    path = tmp_path / "requests.xlsx"
    export_request_sheet(PROFILES, 2026, 10, path)

    write_cell(path, 4, 3, "公")      # 希望休
    write_cell(path, 4, 5, "日③")     # 希望出勤
    write_cell(path, 4, 8, "有")      # 有給
    write_cell(path, 4, 9, "夏")      # 夏休

    requests, warnings = import_request_sheet(path)
    assert warnings == []
    r = requests[0]
    assert r.wish_off_days() == [3]
    assert r.wish_work_days() == {5: "日③"}
    assert r.absence_days() == {8: "有", 9: "夏"}
